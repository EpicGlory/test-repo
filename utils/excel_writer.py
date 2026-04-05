"""Excel output writer with formatting for luxury home data."""

import os
from datetime import datetime
from typing import List

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from models.property import Property


# Column order and display headers
COLUMNS = [
    ("address", "Address"),
    ("city", "City"),
    ("zip_code", "Zip Code"),
    ("price", "Price"),
    ("architect", "Architect"),
    ("designer", "Designer"),
    ("builder", "Builder"),
    ("community", "Community"),
    ("community_subdivision", "Community Subdivision"),
    ("year_built", "Year Built"),
    ("phase", "Phase"),
    ("sqft", "Sq Ft"),
    ("bedrooms", "Bedrooms"),
    ("bathrooms", "Bathrooms"),
    ("lot_size", "Lot Size"),
    ("stories", "Stories"),
    ("style", "Style"),
    ("details", "Details"),
    ("door_details", "door_details"),
    ("listing_url", "Listing URL"),
    ("source", "Source"),
    ("date_scraped", "Date Scraped"),
    ("additional_info", "Additional Info"),
]


def write_to_excel(
    properties: List[Property],
    output_dir: str = "output",
    filename_prefix: str = "luxury_homes_scottsdale_pv",
    update_file: str = None,
) -> str:
    """Write properties to a formatted Excel file.

    Args:
        properties: List of Property objects to write.
        output_dir: Directory for output files.
        filename_prefix: Prefix for the output filename.
        update_file: Path to existing file to update (optional).

    Returns:
        Path to the generated Excel file.
    """
    os.makedirs(output_dir, exist_ok=True)

    # Build DataFrame from properties
    field_names = [col[0] for col in COLUMNS]
    header_names = [col[1] for col in COLUMNS]

    rows = []
    for prop in properties:
        prop_dict = prop.to_dict()
        rows.append([prop_dict.get(field, "") for field in field_names])

    df = pd.DataFrame(rows, columns=header_names)

    # Determine output path
    if update_file and os.path.exists(update_file):
        output_path = update_file
    else:
        date_str = datetime.now().strftime("%Y-%m-%d")
        output_path = os.path.join(output_dir, f"{filename_prefix}_{date_str}.xlsx")

    # Write to Excel
    df.to_excel(output_path, index=False, sheet_name="Luxury Homes", engine="openpyxl")

    # Apply formatting
    _format_workbook(output_path, len(properties))

    return output_path


def _format_workbook(filepath: str, row_count: int):
    """Apply professional formatting to the Excel workbook."""
    wb = load_workbook(filepath)
    ws = wb.active

    # Style definitions
    header_font = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    data_font = Font(name="Calibri", size=10)
    data_alignment = Alignment(vertical="top", wrap_text=True)

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    alt_fill = PatternFill(start_color="F2F7FB", end_color="F2F7FB", fill_type="solid")

    # Format header row
    for col_idx in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # Format data rows
    for row_idx in range(2, row_count + 2):
        for col_idx in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.alignment = data_alignment
            cell.border = thin_border

            # Alternating row colors
            if row_idx % 2 == 0:
                cell.fill = alt_fill

    # Auto-size columns (approximate)
    for col_idx, (field, header) in enumerate(COLUMNS, 1):
        max_len = len(header)

        for row_idx in range(2, min(row_count + 2, 52)):  # Sample first 50 rows
            val = str(ws.cell(row=row_idx, column=col_idx).value or "")
            max_len = max(max_len, min(len(val), 60))

        # Set column width with padding
        col_letter = get_column_letter(col_idx)
        if field == "details":
            ws.column_dimensions[col_letter].width = 50
        elif field in ("listing_url", "additional_info"):
            ws.column_dimensions[col_letter].width = 40
        elif field == "address":
            ws.column_dimensions[col_letter].width = 35
        else:
            ws.column_dimensions[col_letter].width = min(max_len + 4, 30)

    # Freeze header row
    ws.freeze_panes = "A2"

    # Set row height for header
    ws.row_dimensions[1].height = 30

    wb.save(filepath)
